import tempfile
import unittest
from pathlib import Path
from enterprise_doc_ai.extract import extract_invoice, read_pages

GOOD = "Số hóa đơn: HD-2026-0042\nNgày lập: 05/10/2026\nNhà cung cấp: Công ty Demo\nTổng thanh toán: 12.500.000 VNĐ"

class ExtractionTests(unittest.TestCase):
    def test_complete_with_provenance(self):
        result = extract_invoice(GOOD, "fixture.txt", page=2)
        self.assertEqual(result["status"], "complete")
        self.assertEqual(result["fields"]["total_vnd"]["value"], 12500000)
        self.assertEqual(result["fields"]["date"]["value"], "2026-10-05")
        for field in result["fields"].values():
            self.assertEqual(field["page"], 2)
            self.assertEqual(GOOD[field["span"][0]:field["span"][1]].strip(), field["evidence"])

    def test_ambiguous_total_and_invalid_date(self):
        result = extract_invoice(GOOD + "\nTổng thanh toán: 10.000.000", "fixture")
        self.assertIn("ambiguous_total_vnd", result["errors"])
        invalid = extract_invoice(GOOD.replace("05/10/2026", "32/10/2026"), "fixture")
        self.assertIn("invalid_date", invalid["errors"])

    def test_missing_is_not_inferred(self):
        result = extract_invoice("Số hóa đơn: HD-2026-0042", "fixture")
        self.assertEqual(result["status"], "needs_review")
        self.assertIsNone(result["fields"]["date"])
        self.assertIsNone(result["fields"]["total_vnd"])

    def test_invalid_total(self):
        invalid = extract_invoice(GOOD.replace("12.500.000", "12.50.000"), "fixture")
        self.assertIn("invalid_total_format", invalid["errors"])
        zero = extract_invoice(GOOD.replace("12.500.000", "0"), "fixture")
        self.assertIn("nonpositive_total", zero["errors"])

    def test_docx_table_text_is_preserved(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx not installed")
        with tempfile.TemporaryDirectory() as tmp:
            f = Path(tmp) / "invoice.docx"
            doc = Document()
            doc.add_paragraph("Hóa đơn đầu vào")
            row = doc.add_table(rows=1, cols=2).rows[0]
            row.cells[0].text = "Số hóa đơn"
            row.cells[1].text = "HD-2026-0042"
            doc.save(str(f))
            page = read_pages(f)
            self.assertEqual(len(page), 1)
            self.assertIn("HD-2026-0042", page[0][1])

    def test_pdf_blank_page_returns_no_fake_text(self):
        try:
            from pypdf import PdfWriter
        except ImportError:
            self.skipTest("pypdf not installed")
        with tempfile.TemporaryDirectory() as tmp:
            f = Path(tmp) / "scanned.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=200, height=300)
            with f.open("wb") as fp:
                writer.write(fp)
            self.assertEqual(read_pages(f), [(1, "")])

    def test_file_loader_limits(self):
        with tempfile.TemporaryDirectory() as tmp:
            f = Path(tmp) / "invoice.txt"
            f.write_text(GOOD, encoding="utf-8")
            self.assertEqual(read_pages(f), [(1, GOOD)])
            with self.assertRaises(ValueError): read_pages(f, max_bytes=1)
            with self.assertRaises(ValueError): read_pages(f.with_suffix(".exe"))
        with self.assertRaises(ValueError): extract_invoice(GOOD, "fixture", page=0)

if __name__ == "__main__":
    unittest.main()
