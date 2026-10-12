"""Traceable, deterministic Vietnamese invoice extraction without hallucinated fields.

Does not replace OCR or a probabilistic document-understanding model.
"""
from __future__ import annotations
from dataclasses import dataclass
from pathlib import Path
from datetime import datetime
import re

PATTERNS = {
    "invoice_number": re.compile(r"(?im)^\s*(?:Số hóa đơn|Invoice\s*(?:number|no\.?))\s*[:：#]\s*([A-Za-z0-9][\w./-]{1,35})\s*$"),
    "date": re.compile(r"(?im)^\s*(?:Ngày lập|Invoice date)\s*[:：]\s*(\d{2}/\d{2}/\d{4})\s*$"),
    "supplier": re.compile(r"(?im)^\s*(?:Nhà cung cấp|Supplier)\s*[:：]\s*(.{2,100}?)\s*$"),
    "total_vnd": re.compile(r"(?im)^\s*(?:Tổng thanh toán|Total(?: amount)?\s*\(VND\))\s*[:：]\s*([\d., ]+)\s*(?:VNĐ|VND|đ)?\s*$"),
}

def extract_invoice(text: str, source: str, page: int = 1) -> dict:
    if page < 1:
        raise ValueError("page must be >= 1")
    fields: dict[str, dict | None] = {}
    errors: list[str] = []
    for name, pattern in PATTERNS.items():
        matches = list(pattern.finditer(text))
        if len(matches) != 1:
            if len(matches) > 1:
                errors.append(f"ambiguous_{name}")
            fields[name] = None
            continue
        match = matches[0]
        raw = match.group(1).strip()
        value: str | int = raw
        if name == "date":
            try:
                dt = datetime.strptime(raw, "%d/%m/%Y")
                if dt.strftime("%d/%m/%Y") != raw:
                    raise ValueError()
                value = dt.date().isoformat()
            except ValueError:
                errors.append("invalid_date")
                fields[name] = None
                continue
        if name == "total_vnd":
            raw_number = raw.replace(" ", "")
            if not re.fullmatch(r"\d{1,3}(?:\.\d{3})*|\d+", raw_number):
                errors.append("invalid_total_format")
                fields[name] = None
                continue
            value = int(raw_number.replace(".", ""))
            if value <= 0:
                errors.append("nonpositive_total")
                fields[name] = None
                continue
        fields[name] = {"value": value, "evidence": match.group(0).strip(), "source": source,
                        "page": page, "span": [match.start(), match.end()]}
    complete = all(v is not None for v in fields.values()) and not errors
    return {"status": "complete" if complete else "needs_review", "fields": fields,
            "errors": errors, "review_required": not complete, "extraction": "regex_baseline"}

def read_pages(path: str | Path, max_bytes: int = 5_000_000, max_pages: int = 30) -> list[tuple[int, str]]:
    """Local-only ingestion. Text PDF/DOCX/Markdown; scanned PDF may have empty text."""
    path = Path(path)
    if path.is_symlink() or not path.is_file():
        raise ValueError("Only regular non-symlink files are allowed")
    suffix = path.suffix.lower()
    if suffix not in {".md", ".txt", ".pdf", ".docx"}:
        raise ValueError("Unsupported extension")
    if path.stat().st_size > max_bytes:
        raise ValueError("File too large")
    if suffix in {".txt", ".md"}:
        return [(1, path.read_text(encoding="utf-8"))]
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("pip install pypdf") from exc
        reader = PdfReader(str(path))
        if len(reader.pages) > max_pages:
            raise ValueError("Too many pages")
        return [(i + 1, p.extract_text() or "") for i, p in enumerate(reader.pages)]
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:
        raise RuntimeError("pip install python-docx") from exc
    document = DocxDocument(str(path))
    # DOCX has no reliable page count without rendering; page 1 means
    # 'logical document', NOT an accurate original page coordinate.
    content = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            content.append(" | ".join(cell.text for cell in row.cells))
    return [(1, "\n".join(content))]

def ocr_scanned_pdf(path: str | Path, languages: str = "vie+eng", max_pages: int = 10) -> list[tuple[int, str]]:
    """Optional OCR adapter using local Tesseract + PyMuPDF.

    Downloads nothing. Requires installed OCR language packs. Do not use
    on sensitive documents without deployment-specific security review.
    """
    try:
        import fitz
        from PIL import Image
        import pytesseract
    except ImportError as exc:
        raise RuntimeError("pip install pymupdf pillow pytesseract") from exc
    path = Path(path)
    if path.suffix.lower() != ".pdf" or path.is_symlink() or path.stat().st_size > 5_000_000:
        raise ValueError("Only local PDF <= 5 MB accepted")
    results = []
    with fitz.open(path) as document:
        if len(document) > max_pages:
            raise ValueError("Page limit exceeded")
        for i, page in enumerate(document):
            raster = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
            img = Image.frombytes("RGB", [raster.width, raster.height], raster.samples)
            results.append((i + 1, pytesseract.image_to_string(img, lang=languages)))
    return results
