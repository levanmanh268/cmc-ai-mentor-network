"""Offline-first retrieval and grounded answering. No network calls by default."""
from __future__ import annotations
from dataclasses import dataclass
from pathlib import Path
from collections import Counter
import math
import re
from typing import Protocol

TOKEN = re.compile(r"[^\W_]+", flags=re.UNICODE)
SUPPORTED = {".txt", ".md", ".pdf", ".docx"}

def tokenize(text: str) -> list[str]:
    return TOKEN.findall(text.lower())

@dataclass(frozen=True)
class Document:
    id: str
    text: str
    source: str

@dataclass(frozen=True)
class Hit:
    document: Document
    score: float

def chunk_text(text: str, source: str, size: int = 100, overlap: int = 20) -> list[Document]:
    if size < 1 or overlap < 0 or overlap >= size:
        raise ValueError("Require size > overlap >= 0")
    words = text.split()
    step = size - overlap
    return [
        Document(f"{source}#chunk-{start // step + 1}", " ".join(words[start:start + size]), source)
        for start in range(0, len(words), step)
        if words[start:start + size]
    ]

def load_document(path: str | Path, max_bytes: int = 3_000_000) -> str:
    path = Path(path)
    if path.suffix.lower() not in SUPPORTED:
        raise ValueError("Allowed file types: .txt, .md, .pdf, .docx")
    if path.stat().st_size > max_bytes:
        raise ValueError("File exceeds 3 MB safety limit")
    if path.suffix.lower() in {".md", ".txt"}:
        return path.read_text(encoding="utf-8")
    if path.suffix.lower() == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("Install pypdf to parse PDF documents") from exc
        return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    try:
        from docx import Document as WordDocument
    except ImportError as exc:
        raise RuntimeError("Install python-docx to parse DOCX documents") from exc
    return "\n".join(p.text for p in WordDocument(str(path)).paragraphs)

class BM25Index:
    """Okapi BM25 with deterministic tie breaks; simple, auditable sparse baseline."""
    def __init__(self, documents: list[Document], k1: float = 1.5, b: float = 0.75):
        self.documents = list(documents)
        self.k1, self.b = k1, b
        self.tf = [Counter(tokenize(d.text)) for d in self.documents]
        self.df = Counter(t for counts in self.tf for t in counts)
        self.avglen = sum(sum(x.values()) for x in self.tf) / max(len(self.tf), 1)

    def search(self, question: str, k: int = 3) -> list[Hit]:
        if k < 1:
            return []
        query = set(tokenize(question))
        n = len(self.documents)
        ranked = []
        for doc, counts in zip(self.documents, self.tf):
            length = sum(counts.values())
            score = 0.0
            for term in query:
                if term not in counts:
                    continue
                idf = math.log(1 + (n - self.df[term] + 0.5) / (self.df[term] + 0.5))
                score += idf * counts[term] * (self.k1 + 1) / (counts[term] + self.k1 * (1 - self.b + self.b * length / max(self.avglen, 1)))
            if score > 0:
                ranked.append(Hit(doc, score))
        return sorted(ranked, key=lambda x: (-x.score, x.document.id))[:k]

class TFIDFIndex:
    """L2-normalized TF-IDF baseline for explicit comparison with BM25."""
    def __init__(self, documents: list[Document]):
        self.documents = list(documents)
        self.df = Counter(t for d in documents for t in set(tokenize(d.text)))
        self.idf = {t: math.log((1 + len(documents)) / (1 + df)) + 1 for t, df in self.df.items()}
        self.vectors = [self._vector(d.text) for d in documents]

    def _vector(self, text: str) -> dict[str, float]:
        counts = Counter(tokenize(text))
        vec = {term: (1 + math.log(freq)) * self.idf[term] for term, freq in counts.items() if term in self.idf}
        norm = math.sqrt(sum(v * v for v in vec.values())) or 1
        return {term: value / norm for term, value in vec.items()}

    def search(self, question: str, k: int = 3) -> list[Hit]:
        if k < 1:
            return []
        query = self._vector(question)
        hits = [Hit(d, sum(value * vec.get(term, 0.0) for term, value in query.items())) for d, vec in zip(self.documents, self.vectors)]
        return sorted((h for h in hits if h.score > 0), key=lambda h: (-h.score, h.document.id))[:k]

class Encoder(Protocol):
    def encode(self, texts: list[str]) -> list[list[float]]: ...

class SentenceTransformerEncoder:
    """Optional real semantic embeddings. Download model only when explicitly instantiated."""
    def __init__(self, model_name: str = "sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2"):
        try:
            from sentence_transformers import SentenceTransformer
        except ImportError as exc:
            raise RuntimeError("pip install sentence-transformers") from exc
        self.model = SentenceTransformer(model_name)

    def encode(self, texts: list[str]) -> list[list[float]]:
        return self.model.encode(texts, normalize_embeddings=True).tolist()

class DenseIndex:
    """Dependency-injected embedding interface; cosine normalization done locally."""
    def __init__(self, documents: list[Document], encoder: Encoder):
        self.documents = list(documents)
        self.encoder = encoder
        self.vectors = self._normalized(encoder.encode([d.text for d in documents]))

    @staticmethod
    def _normalized(vectors: list[list[float]]) -> list[list[float]]:
        out = []
        for vec in vectors:
            norm = math.sqrt(sum(float(x) ** 2 for x in vec)) or 1.0
            out.append([float(x) / norm for x in vec])
        return out

    def search(self, question: str, k: int = 3) -> list[Hit]:
        if not self.documents or k < 1:
            return []
        query = self._normalized(self.encoder.encode([question]))[0]
        if len(query) != len(self.vectors[0]):
            raise ValueError("Embedding dimensions do not match")
        hits = [Hit(doc, sum(a * b for a, b in zip(query, vec))) for doc, vec in zip(self.documents, self.vectors)]
        return sorted((h for h in hits if h.score > 0), key=lambda h: (-h.score, h.document.id))[:k]

class GroundedAssistant:
    def __init__(self, index: BM25Index | TFIDFIndex | DenseIndex, min_score: float = 0.05):
        self.index, self.min_score = index, min_score

    def answer(self, question: str, k: int = 3) -> dict:
        hits = [h for h in self.index.search(question, k) if h.score >= self.min_score]
        if not hits:
            return {"status": "abstained", "answer": "Không tìm thấy thông tin đủ liên quan trong tài liệu.", "citations": [], "mode": "extractive"}
        citations = [{"id": h.document.id, "source": h.document.source, "score": round(h.score, 6)} for h in hits]
        # Explicit extractive baseline: never claim this as an LLM-generated answer.
        excerpt = hits[0].document.text[:500].strip()
        return {"status": "answered", "answer": excerpt, "citations": citations, "mode": "extractive"}

class SafeAgent:
    """Allowlisted, deterministic tool routing. No shell, network, or write operations."""
    def __init__(self, assistant: GroundedAssistant):
        self.assistant = assistant

    def run(self, tool: str, question: str = "") -> dict:
        if tool == "search_knowledge":
            return {"tool": tool, "result": self.assistant.answer(question)}
        if tool == "list_sources":
            sources = sorted({d.source for d in self.assistant.index.documents})
            return {"tool": tool, "result": sources}
        return {"tool": tool, "error": "Tool not allowed", "allowed": ["search_knowledge", "list_sources"]}
